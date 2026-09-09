"""Parse a TCE catalogue Excel workbook into the app data structure.
Mirrors the logic used to build catalogue_data.json so uploaded files
with the same structure can refresh the catalogue."""
import io
import re
from datetime import datetime
import openpyxl


def _clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 4)
    if isinstance(v, str):
        v = v.replace("\u202f", "").replace(" ", "").replace(",", ".").strip()
        try:
            return round(float(v), 4)
        except ValueError:
            return None
    return None


def parse_catalogue(content: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    lot_sheets = [s for s in wb.sheetnames if re.match(r"^\d", s)]
    lot_sheets = [s for s in lot_sheets if not s.startswith("00")]

    articles = []
    for sheet in lot_sheets:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        for r in rows[1:]:
            code = _clean(r[0]) if len(r) > 0 else None
            if not code or not re.match(r"^[A-Z]{2,4}-\d", str(code)):
                continue
            a = {
                "code": code,
                "lot": sheet,
                "sous_famille": _clean(r[1]) if len(r) > 1 else None,
                "article": _clean(r[2]) if len(r) > 2 else None,
                "unite": _clean(r[3]) if len(r) > 3 else None,
                "designation_fournisseur": _clean(r[4]) if len(r) > 4 else None,
                "marque": _clean(r[5]) if len(r) > 5 else None,
                "ref_fabricant": _clean(r[6]) if len(r) > 6 else None,
                "ref_fournisseur": _clean(r[7]) if len(r) > 7 else None,
                "fournisseur_retenu": _clean(r[8]) if len(r) > 8 else None,
                "prix_achat_ht": _num(r[9]) if len(r) > 9 else None,
                "type_prix": _clean(r[10]) if len(r) > 10 else None,
                "prix_public_ref": _num(r[11]) if len(r) > 11 else None,
                "remise": _num(r[12]) if len(r) > 12 else None,
                "date_prix": _clean(r[13]) if len(r) > 13 else None,
                "fiabilite": _clean(r[14]) if len(r) > 14 else None,
                "conditionnement": _clean(r[15]) if len(r) > 15 else None,
                "point_controle": _clean(r[16]) if len(r) > 16 else None,
                "marge_pct": _num(r[17]) if len(r) > 17 else None,
                "prix_vente_ht": _num(r[18]) if len(r) > 18 else None,
                "tva_pct": _num(r[19]) if len(r) > 19 else 20,
                "retrait": _clean(r[20]) if len(r) > 20 else None,
                "fiche_produit": _clean(r[21]) if len(r) > 21 else None,
                "offers": [],
            }
            if isinstance(a["date_prix"], datetime):
                a["date_prix"] = a["date_prix"].strftime("%d/%m/%Y")
            articles.append(a)

    by_code = {a["code"]: a for a in articles}

    if "Base articles" in wb.sheetnames:
        ws = wb["Base articles"]
        for row in ws.iter_rows(min_row=2):
            r = [c.value for c in row]
            code = _clean(r[0]) if len(r) > 0 else None
            if not code or code not in by_code:
                continue
            fiche_url = None
            if len(row) > 17 and row[17] is not None:
                hl = row[17].hyperlink
                if hl and hl.target:
                    fiche_url = hl.target
                elif isinstance(row[17].value, str) and row[17].value.startswith("http"):
                    fiche_url = row[17].value
            offer = {
                "fournisseur": _clean(r[5]),
                "type_prix": _clean(r[6]),
                "prix_ht": _num(r[7]),
                "prix_public": _num(r[8]),
                "remise": _num(r[9]),
                "designation": _clean(r[10]),
                "marque": _clean(r[11]),
                "ref_fournisseur": _clean(r[12]),
                "unite_vente": _clean(r[13]),
                "conditionnement": _clean(r[14]),
                "retrait": _clean(r[15]),
                "fiabilite": _clean(r[16]),
                "fiche_produit": fiche_url,
                "rang": _num(r[19]),
            }
            if offer["fournisseur"] and offer["prix_ht"] is not None:
                by_code[code]["offers"].append(offer)

    def parse_simple(sheet, start, mapping):
        if sheet not in wb.sheetnames:
            return []
        out = []
        for r in list(wb[sheet].iter_rows(values_only=True))[start:]:
            first = _clean(r[0]) if len(r) > 0 else None
            if not first:
                continue
            rec = {}
            for key, (idx, kind) in mapping.items():
                val = r[idx] if len(r) > idx else None
                rec[key] = _num(val) if kind == "num" else _clean(val)
            out.append(rec)
        return out

    synthese = []
    if "Synthèse" in wb.sheetnames:
        for r in list(wb["Synthèse"].iter_rows(values_only=True))[4:]:
            lot = _clean(r[0])
            if not lot or not re.match(r"^\d", str(lot)):
                continue
            synthese.append({
                "lot": lot, "articles": _num(r[1]), "prix_releves": _num(r[2]),
                "dont_votre_tarif": _num(r[3]), "comparables": _num(r[4]),
                "a_completer": _num(r[5]), "couverture": _num(r[6]), "prix_moyen_ht": _num(r[7]),
            })

    sources = parse_simple("Sources", 4, {
        "enseigne": (0, "str"), "site": (1, "str"), "acces": (2, "str"),
        "statut": (3, "str"), "tire": (4, "str"),
        "produits_examines": (5, "num"), "offres_retenues": (6, "num"),
    })
    parametres = parse_simple("Paramètres", 4, {
        "parametre": (0, "str"), "valeur": (1, "str"), "commentaire": (2, "str"),
    })
    controls = parse_simple("Contrôles qualité", 4, {
        "lot": (0, "str"), "article": (1, "str"), "fournisseur": (2, "str"),
        "prix_ht": (3, "num"), "nature": (4, "str"), "action": (5, "str"),
    })

    return {
        "articles": articles,
        "synthese": synthese,
        "sources": sources,
        "parametres": parametres,
        "controls": controls,
    }
