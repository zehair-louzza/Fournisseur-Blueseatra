import openpyxl
wb = openpyxl.load_workbook('/app/catalogue.xlsx', data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print("\n" + "="*80)
    print(f"SHEET: {ws.title}  dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 12:
            break
        cells = [str(c)[:22] if c is not None else "" for c in row]
        print(f"r{i}:", " | ".join(cells))
